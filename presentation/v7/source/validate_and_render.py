from pathlib import Path
from zipfile import ZipFile
from PIL import Image,ImageDraw
import fitz,json,re,xml.etree.ElementTree as E
from pptx import Presentation
b=Path(__file__).resolve().parents[1];out=b/'slides';out.mkdir(exist_ok=True)
r=Presentation(b/'VineWeave-Canopy-v7.pptx');assert len(r.slides)==18
ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
with ZipFile(b/'VineWeave-Canopy-v7.pptx') as z:
 assert z.testzip() is None
 charts=[x for x in z.namelist() if re.fullmatch(r'ppt/charts/chart\d+\.xml',x)];workbooks=[x for x in z.namelist() if x.endswith('.xlsx')];gifs=[x for x in z.namelist() if x.endswith('.gif')]
 print('Native charts:',len(charts),'workbooks:',len(workbooks),'GIFs:',len(gifs))
 assert len(charts)==5 and len(workbooks)==5 and len(gifs)==2
 for name in gifs:
  import io
  im=Image.open(io.BytesIO(z.read(name)));print(name,'retained frames',im.n_frames)
 for i,sl in enumerate(r.slides,1):
  assert sl.has_notes_slide
  assert sum(sh.has_text_frame for sh in sl.shapes)>3
  txt=' '.join(sh.text for sh in sl.shapes if sh.has_text_frame)
  assert not any(x in txt for x in ['D_BASE','D版','D canopy','D geometry'])
  for sh in sl.shapes:
   if sh.width>=0 and sh.height>=0:
    assert sh.left>=-1000 and sh.top>=-1000 and sh.left+sh.width<=r.slide_width+1000 and sh.top+sh.height<=r.slide_height+1000,(i,sh.name,sh.left,sh.top,sh.width,sh.height)
 print('18 slides: notes, editable text and slide bounds passed.')
d=fitz.open(b/'VineWeave-Canopy-v7.pdf');assert len(d)==18
issues=[]
for i,page in enumerate(d):
 for block in page.get_text('blocks'):
  if block[0]<-1 or block[1]<-1 or block[2]>page.rect.width+1 or block[3]>page.rect.height+1:issues.append((i+1,'pdf text outside page',block[:4]))
 pix=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False);pix.save(out/f'slide-{i+1:02d}.png')
for start in range(0,18,4):
 batch=list(range(start,min(start+4,18)));sheet=Image.new('RGB',(1600,960),'#DDE0D8');dr=ImageDraw.Draw(sheet)
 for j,i in enumerate(batch):
  im=Image.open(out/f'slide-{i+1:02d}.png');im.thumbnail((780,440));x=(j%2)*800+10;y=(j//2)*480+28
  sheet.paste(im,(x,y));dr.text((x,y-20),f'SLIDE {i+1:02d}',fill='black')
 sheet.save(b/f'contact-{start//4+1}.jpg')
print('PDF text-bound issues:',issues)
(b/'validation-machine.json').write_text(json.dumps({'slides':len(r.slides),'native_charts':len(charts),'embedded_workbooks':len(workbooks),'gifs':len(gifs),'pdf_text_bounds_issues':issues,'powerpoint_playback_tested':False},indent=2))

for sl in r.slides:
 texts=[]
 for sh in sl.shapes:
  if sh.has_text_frame:texts.append(sh.text)
  if sh.has_table:texts.extend(c.text for row in sh.table.rows for c in row.cells)
 texts.append(sl.notes_slide.notes_text_frame.text)
 txt=' '.join(texts).lower()
 assert not any(term in txt for term in ['growing seasons','route time','maturity date','growth-duration','p6 plant','1 / 2 / 3 m','1.5 m assumed','calibrate plant establishment'])
print('Removed plant-time content: text, tables and notes passed.')

assert not any('Translate the brief' in sh.text for sl in r.slides for sh in sl.shapes if sh.has_text_frame)
