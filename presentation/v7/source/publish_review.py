from pathlib import Path
from pptx import Presentation
from docx import Document
from docx.shared import Pt
import re,json,zipfile
b=Path(__file__).resolve().parents[1];notes=json.loads((b/'speaker-notes.json').read_text());records=json.loads((b/'slide-records.json').read_text());basis=json.loads((b/'source/source-basis.json').read_text())
assert len(notes)==18
words=sum(len(re.findall(r"[\w’'-]+",n['script']+' '+n['transition'])) for n in notes[:-1])
intro=f'{words} spoken English words on slides 1–17. At 120–135 words/minute this is {words/135:.1f}–{words/120:.1f} minutes of continuous speech, before pauses and slide changes. Rehearse with figures and reserve Q&A within 10–12 minutes. Takeaways, Q&A and source-basis sections are not extra timed speech.'
D=Document();D.styles['Normal'].font.name='Noto Sans';D.styles['Normal'].font.size=Pt(11);D.add_heading('VineWeave Canopy — Presenter Guide V7',0);D.add_paragraph(intro);lines=['# Presenter Guide V7','',intro,'']
hints={4:'先讲材料分区与数量；这是材料方案，不是整体安全结论。',5:'表中单价是按相同全长比较材料输入，不是最终单元总价；材料常数仍为工作簿输入。',6:'按原工作簿解释钢支撑与竹冠面；别把竹材单项指标的优势说成整套造价更低。',7:'先指33.25mm目标，再看两种理想边界；原表明确要求节点详图与转动刚度验算。',11:'先分清两类指标，再读后面的百分比；不能混用、相加或换算摄氏度。',12:'正午98.7%不能代表整段使用时间；时段与座椅位置要一起考虑。',13:'只说原CSV给出的参数情景结果；同覆盖比例不等于相同种植面积和材料投入。',14:'80到95覆盖增加2.8个百分点是时段均值；相同50%分布比较是正午值，别混成同一组。'}
for n in notes:
 title=f"{n['slide']:02d} | {n['title']}";D.add_heading(title,1);lines+=['## '+title,'']
 if n['slide'] in hints:D.add_paragraph('中文提示：'+hints[n['slide']]);lines+=['中文提示：'+hints[n['slide']],'']
 for label,key in [('Takeaway','takeaway'),('Spoken script','script'),('Transition','transition'),('Q&A','qa')]:D.add_paragraph(label+': '+n[key]);lines+=[label+': '+n[key],'']
 D.add_heading('Source basis — not spoken',2);lines+=['Source basis — not spoken:','']
 for src in n['sourceBasis']['sources']:D.add_paragraph(src);lines+=[src,'']
 D.add_paragraph(n['sourceBasis']['basis']);lines+=[n['sourceBasis']['basis'],'']
D.save(b/'Presenter-Guide-v7.docx');(b/'Presenter-Guide-v7.md').write_text('\n'.join(lines))
preview=['# VineWeave Canopy — V7','','**18 slides including references.** [Editable PowerPoint](VineWeave-Canopy-v7.pptx) · [PDF](VineWeave-Canopy-v7.pdf) · [Presenter guide](Presenter-Guide-v7.md) · [Statement/source audit](Statement-Evidence.md)','','Material reasoning: pages 4–6. Node/deflection: page 7. Shade methods/results: pages 11–14. No plant-growth timing or course-mapping page.','']
for r in records:preview += [f"## {r['number']:02d} — {r['title']}",'',f"![Slide {r['number']:02d}](slides/slide-{r['number']:02d}.png)",'']
(b/'Preview.md').write_text('\n'.join(preview))
# Index every native text statement, table cell, chart entry and presenter sentence
# to the corresponding supplied-source section/cells. Evidence groups and limits
# remain explicit; this is traceability, not independent certification of inputs.
r=Presentation(b/'VineWeave-Canopy-v7.pptx');audit=['# V7逐项文字与依据对照','','范围：主页面原生文字、表格单元、图表数值及讲稿句子。每页先列明来自用户材料的出处与依据，再列该页表述；同页表述对应其下所列来源。问句、标题和过渡用于组织已列材料，不构成新增实验结果。','','“有材料出处”不等于材料中的假设已获验证。工程竹常数、湿植被荷载、透射率、报价和节点约束仍按输入/假设标注。原报告列出的外部文献并未冒充本轮已核验全文。','','修改依据：Cost-Analysis.xlsx、原分析报告、参数/覆盖补充报告与CSV；新增算术复核见[Claim-Checks.json](Claim-Checks.json)。没有新增生长速率、实测温降、风速或结构安全结论。',''];all_statements=[]
for i,sl in enumerate(r.slides,1):
 ref=basis[str(i)];audit += [f"## {i:02d} — {records[i-1]['title']}",'','出处：','']+['- '+x for x in ref['sources']]+['','材料依据／限定：'+ref['basis'],'','### 页面表述','']
 statements=[]
 for sh in sl.shapes:
  if sh.has_text_frame:
   for para in sh.text_frame.paragraphs:
    if para.text.strip():statements.append(('页面文字',para.text))
  if sh.has_table:
   for row in sh.table.rows:
    for c in row.cells:
     if c.text.strip():statements.append(('表格',c.text))
  if sh.has_chart:
   for ser in sh.chart.series:
    statements += [('图表数据',f'{cat.label}: {val}') for cat,val in zip(sh.chart.plots[0].categories,ser.values)]
 for kind,text in statements:
  for sentence in re.split(r'(?<=[.!?])\s+(?=[A-Z])',text):
   audit += ['- '+kind+'：'+sentence.replace('\n',' / ')];all_statements.append({'slide':i,'location':kind,'statement':sentence,'basis':str(i)})
 audit += ['','### 讲稿及答辩表述','']
 n=notes[i-1]
 for field in ['takeaway','script','transition','qa']:
  for sentence in re.split(r'(?<=[.!?])\s+(?=[A-Z])',n[field]):audit+=['- '+field+'：'+sentence];all_statements.append({'slide':i,'location':field,'statement':sentence,'basis':str(i)})
 audit+=['']
(b/'Statement-Evidence.md').write_text('\n'.join(audit));(b/'Statement-Evidence.json').write_text(json.dumps({'basis':basis,'statements':all_statements},ensure_ascii=False,indent=2))
(b/'README.md').write_text('''# VineWeave Canopy — V7

**18 English slides including references.** Material choice, node/deflection explanation and shade interpretation revised using the supplied project materials.

- **[View every slide](Preview.md)**
- [Editable PowerPoint](VineWeave-Canopy-v7.pptx) · [Static PDF](VineWeave-Canopy-v7.pdf)
- [Presenter guide with source basis](Presenter-Guide-v7.md) · [Word guide](Presenter-Guide-v7.docx)
- [Presentation bundle](Presentation-v7.zip)
- [Statement/source audit](Statement-Evidence.md) · [Numeric claim checks](Claim-Checks.json)
- [Changes and evidence status](Changes-and-Evidence.md) · [Validation](Validation.md)

![Materials and node pages](contact-2.jpg)

One additional material-decision page separates the numeric screening table from the source’s hybrid-selection rationale. The node page now compares the two ideal beam cases with the adopted L/200 target. The shade-method page distinguishes continuous-envelope overlap from equivalent direct-beam blockage; results state their metric/time window and identify source-supported implications.

No new measurements, physical simulations, product evidence, plant-growth timing or engineering certification are claimed. The course-mapping page stays removed. Text, tables, analytical charts and mechanism shapes remain editable; model renders and GIFs remain images. Five charts retain embedded workbooks, and two GIFs remain on pages 14–15. Static gallery/PDF use keyframes.

Source references are embedded in all slide notes and repeated in the statement audit. Source assumptions are kept distinct from observations and arithmetic derived from the data. Original private attachments are not republished. Companion schematic drawings and schedules remain in [V3 deliverables](../v3/submission/).
''')
(b/'Changes-and-Evidence.md').write_text('''# V7改动与证据范围

| 页面 | 修改 | 来自用户材料的依据 |
| --- | --- | --- |
| 4–5 | 保留材料分区、数量及原材料筛选数据；明确表格费用不是整套单元造价 | Cost-Analysis.xlsx 材料比选、几何与用量、选型决策矩阵 |
| 6 | 新增一页解释“钢支撑＋工程竹冠层”的既有选择逻辑；把五项准则与前提说清 | 材料比选17–18行、挠度校核46行、选型决策矩阵2/28–34行 |
| 7 | 图中加入原表采用的33.25mm目标，与100.53/20.11mm理想结果并列；明确节点验证要求 | 挠度校核31/32/46行；新增3.02/0.60为原值除以目标的算术 |
| 11 | 两类遮荫指标改用文字问题解释，公式保留在讲稿；强调不能混用 | 原分析报告与补充报告§01–02的定义及公式 |
| 12 | 从正午峰值转向全时段解读，补足晨午后使用判断 | 原CSV及原报告应用/评价章节 |
| 13 | 冠幅标题直接陈述原CSV的52.3%→66.1%，不写普遍最优或结构安全 | 补充包汇总CSV；同覆盖比例不等于同投入的原报告说明 |
| 14 | 区分覆盖率与等效遮挡；区分时段均值与同覆盖正午分布结果 | 原CSV；80→95覆盖的2.8个百分点为原数据相减 |

本轮也清理了此前植物页没有来自已读用户材料的“暖地”表述，改为原植物表中的搭配角色、缠绕方式及原报告要求的本地适生/附着检查。未引用未访问成功的网站。

资料冲突处理：工作簿有旧的铰接/完全可逆说法，也有后续灌浆节点修订；采用后者，并保留第46行明确要求的节点验算前提。不把源表按假定给出的“通过”扩大为真实节点或整体结构安全已通过。

所有原始参数、价格和碳因子仍按其证据状态展示。逐项出处、原文要点和文字清单见Statement-Evidence.md；该清单证明可追溯，不代替产品、标准或论文原文验证。
''')
(b/'Validation.md').write_text(f'''# V7 validation

- 18-slide PPTX package integrity passed; native text/table objects, shape bounds and speaker notes checked.
- Five native charts and five embedded workbooks retained; two GIFs retain 5/25 frames.
- PDF has 18 pages. Text page bounds checked; changed pages visually reviewed.
- Source workbook independently reproduces both beam values with its E/I/q/L inputs. Target and ratios checked.
- Seven parameter means, five coverage means and same-cover noon values match supplied CSVs. The 80→95% cover gain rounds to 2.8 percentage points.
- Source-basis references embedded in all 18 slide notes; statement index covers native text, table cells, chart entries and presenter sentences.
- Plant-growth timing and course-mapping page remain absent. No new physical/biological simulation or external literature verification performed.
- Playback and timing require rehearsal. Source assumptions, prices and node behaviour are not certified.

{intro}
''')
with zipfile.ZipFile(b/'Presentation-v7.zip','w',zipfile.ZIP_DEFLATED) as z:
 for name in ['VineWeave-Canopy-v7.pptx','VineWeave-Canopy-v7.pdf','Presenter-Guide-v7.docx','README.md','Statement-Evidence.md','Claim-Checks.json','Changes-and-Evidence.md','Validation.md']:z.write(b/name,name)
print('Published:',len(notes),'slides;',words,'spoken words;',len(all_statements),'indexed text/data items.')
