from pathlib import Path
import sys,os,json,math
os.environ['MPLCONFIGDIR']='/workspace/ppt-tools/cache/matplotlib';sys.path.insert(0,'/workspace/ppt-tools/model-inspect')
import numpy as np,rhino3dm as rh
import matplotlib;matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import Arc,Rectangle,Circle
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.workbook.properties import CalcProperties
ROOT=Path(__file__).resolve().parents[1];A=ROOT/'assets';O=ROOT/'submission';O.mkdir(exist_ok=True)
B=Path('/workspace/ppt-extension/06_参数与植物生长情景模拟');model=rh.File3dm.Read(str(B/'模型/D_BASE.3dm'));mesh=np.load(B/'复算/input_data/D_BASE.npz');vs=mesh['vertices_m'];curves=[]
for obj in model.Objects:
 if isinstance(obj.Geometry,rh.Curve):
  c=obj.Geometry;pts=np.array([[q.X,q.Y,q.Z]for q in[c.PointAt(t)for t in np.linspace(c.Domain.T0,c.Domain.T1,220)]])/1000;curves.append(pts)
assert len(curves)==59
geo=json.loads(Path('/workspace/temp/research/supplemental-geometric-check.json').read_text());(A/'geometry-definitions.json').write_text(json.dumps(geo,indent=2))
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'pdf.fonttype':42,'ps.fonttype':42})
green='#346A55';navy='#173C51';gold='#AA7E4A'
meta={'outer':{'xmin':float(vs[:,0].min()),'xmax':float(vs[:,0].max()),'ymin':float(vs[:,1].min()),'ymax':float(vs[:,1].max())}}
for key,figsize,lims,cols in [('plan',(6,6),((-4.3,4.3),(-4.7,3.9)),(0,1)),('elevation',(8.6,5.2),((-4.3,4.3),(-.6,4.6)),(0,2))]:
 fig=plt.figure(figsize=figsize);ax=fig.add_axes([0,0,1,1]);
 for c in curves:ax.plot(c[:,cols[0]],c[:,cols[1]],color=green,lw=.75)
 ax.set_xlim(*lims[0]);ax.set_ylim(*lims[1]);ax.set_aspect('equal');ax.axis('off');fig.savefig(A/f'technical-{key}.png',dpi=160,transparent=True);plt.close(fig)
 from PIL import Image
 im=Image.open(A/f'technical-{key}.png');meta[key]={'width':im.width,'height':im.height,'xlim':lims[0],'ylim':lims[1]}
(A/'drawing-meta.json').write_text(json.dumps(meta,indent=2))
# Vector PDF: general arrangement and a separate, clearly schematic interface sheet.
def dim(ax,a,b,text,offset=.0):
 ax.annotate('',xy=b,xytext=a,arrowprops=dict(arrowstyle='|-|',color=gold,lw=.9));mid=(np.array(a)+np.array(b))/2
 ax.text(mid[0],mid[1]+offset,text,color=navy,fontsize=11,ha='center',va='center',bbox=dict(fc='white',ec='none',pad=1))
with PdfPages(O/'Technical-Drawings.pdf') as pdf:
 fig=plt.figure(figsize=(16.54,11.69));fig.text(.045,.95,'VineWeave Canopy | General Arrangement',fontsize=22,color=navy,weight='bold');fig.text(.045,.915,'green building · Group 12 · Schematic design drawing · Dimensions in metres unless noted',fontsize=11,color=green)
 ax=fig.add_axes([.07,.19,.43,.66]);
 for c in curves:ax.plot(c[:,0],c[:,1],color=green,lw=.65)
 q=meta['outer'];ax.set(xlim=(-4.3,4.3),ylim=(-4.7,3.9),aspect='equal');ax.axis('off');dim(ax,(q['xmin'],q['ymin']-.38),(q['xmax'],q['ymin']-.38),'6.22 m',offset=-.22)
 ax.annotate('',xy=(q['xmin']-.4,q['ymax']),xytext=(q['xmin']-.4,q['ymin']),arrowprops=dict(arrowstyle='|-|',color=gold,lw=.9));ax.text(q['xmin']-.63,(q['ymin']+q['ymax'])/2,'6.66 m',rotation=90,va='center',ha='center',fontsize=11,color=navy)
 pair=np.array(geo['pair_xy_m']);ax.plot(pair[:,0],pair[:,1],ls='--',lw=.9,color='#3E6E9E');ax.text(.0,.3,'Max. projected exterior span ≈6.91 m',rotation=49,fontsize=9,color=navy,ha='center',bbox=dict(fc='white',ec='none',alpha=.85,pad=1));ax.set_title('PLAN | Centre-line representation',fontsize=13,color=navy,pad=10)
 ax2=fig.add_axes([.56,.53,.39,.29]);
 for c in curves:ax2.plot(c[:,0],c[:,2],color=green,lw=.65)
 ax2.set(xlim=(-4.3,4.3),ylim=(-.6,4.6),aspect='equal');ax2.axis('off');ax2.plot([-3.6,3.6],[0,0],color='#7E8A86',lw=.6);dim(ax2,(q['xmax']+.38,0),(q['xmax']+.38,3.48),'3.48 m')
 ax2.set_title('ELEVATION | Projected X–Z view',fontsize=13,color=navy,pad=8)
 aa=fig.add_axes([.59,.19,.20,.25]);c=next(c for c in curves if c[:,2].min()<.1);lo=c[-1];hi=c[0];u=(hi[:2]-lo[:2]);u=u/np.linalg.norm(u);xx=(c[:,:2]-lo[:2])@u;zz=c[:,2]-lo[2];hh=float(xx[0]);rise=float(zz[0]);theta=math.degrees(math.atan2(rise,hh));aa.plot(xx,zz,color=green,lw=1);aa.plot([0,hh],[0,rise],'--',color=gold,lw=1);aa.plot([0,hh],[0,0],color='#88958B',lw=.8);aa.add_patch(Arc((0,0),1.4,1.4,theta1=0,theta2=theta,color=gold,lw=1));aa.text(.75,.26,'θ',fontsize=14,color=navy);aa.set(xlim=(-.4,4),ylim=(-.4,4),aspect='equal');aa.axis('off');aa.set_title('INCLINATION DEFINITION',fontsize=11,color=navy)
 fig.text(.80,.38,'12 support endpoint chords:\nθ = 45.9°–50.7° above horizontal.\nThe curved member does not\nhave a single local inclination.',fontsize=10,color=navy,va='top')
 fig.text(.56,.13,'59 members · exterior diameter 80 mm\nSelected crown controls: corner 0.85, expansion 1.08',fontsize=11,color=green)
 fig.text(.045,.065,'Source: supplied model centre lines and exterior mesh; dimensions rounded. Max span uses XY exterior-mesh distances; angles use 3D endpoint chords.\nThis is a general-arrangement drawing, not a fabrication, foundation, structural-safety or accessibility certification.',fontsize=9.5,color='#52615B')
 fig.text(.94,.035,'GA-01',fontsize=10,ha='right',color=green);pdf.savefig(fig);plt.close(fig)
 fig=plt.figure(figsize=(16.54,11.69));fig.text(.045,.95,'VineWeave Canopy | Member Schedule & Interfaces',fontsize=22,color=navy,weight='bold');fig.text(.045,.915,'Proposed build-up · Interface schematic, not an engineered node detail',fontsize=11,color=green)
 ax=fig.add_axes([.05,.43,.53,.39]);ax.axis('off');rows=[['Crown','47','Engineered bamboo, solid Ø80 mm','219.28 m'],['Supports','12','Q355B tube, Ø80×4 mm, galvanised','69.38 m'],['Head nodes','12','Plate + sleeve + adhesive + bolts','Verify rotational stiffness'],['Base','1 set','Flange/anchorage + precast weight','Foundation checks required'],['Planting','2 + 1','Boston ivy + Rangoon creeper','Local suitability pending']];tab=ax.table(cellText=rows,colLabels=['Zone','Qty','Proposed specification','Quantity / condition'],colWidths=[.16,.08,.40,.36],loc='center',cellLoc='left');tab.auto_set_font_size(False);tab.set_fontsize(10.5);tab.scale(1,2.2)
 for (i,j),cell in tab.get_celld().items():cell.set_edgecolor('#CDD6CC');cell.set_facecolor('#EDF1E9'if i==0 else'white')
 na=fig.add_axes([.64,.43,.31,.38]);na.axis('off');na.plot([.2,.2],[.1,.6],color='#3E6E9E',lw=15);na.plot([.2,.82],[.6,.84],color=gold,lw=12);na.add_patch(Rectangle((.13,.54),.15,.13,fc='#C8D7DF',ec=navy));na.add_patch(Circle((.2,.60),.025,fc='white',ec=navy));na.text(.04,.01,'Steel support',fontsize=11,color=navy);na.text(.43,.90,'Bamboo crown',fontsize=11,color=navy);na.text(.34,.44,'Plate / sleeve / adhesive / bolts\nSizing and joint rotation to be verified',fontsize=10,color=navy);na.set(xlim=(0,1),ylim=(0,1))
 fig.text(.06,.32,'ASSEMBLY LOGIC',fontsize=13,color=green,weight='bold');fig.text(.06,.275,'Prefabricate & number  →  Position base  →  Align members  →  Assemble, grout & cure  →  Inspect, plant & restore',fontsize=12,color=navy)
 fig.text(.06,.19,'14 person-days is a preliminary labour allowance, not 14 calendar days.\nLong members and the ≈9.07 t reference counterweight constrain transport/lifting; wet joints limit lossless reversal.',fontsize=11,color=navy)
 fig.text(.045,.065,'Source: Cost-Analysis.xlsx and project build-up. Product data, attachment, loads, global stability, joint restraint and foundation/anchorage remain unverified.\nNo engineering authority or supplier quotation is implied by this schematic.',fontsize=9.5,color='#52615B');fig.text(.94,.035,'GA-02',fontsize=10,ha='right',color=green);pdf.savefig(fig);plt.close(fig)
# One-page group brief, not the original performance-only module.
DOC=Document();sec=DOC.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.5);sec.bottom_margin=Cm(1.5);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
style=DOC.styles['Normal'];style.font.name='Noto Sans';style.font.size=Pt(10);style.paragraph_format.space_after=Pt(4);style.paragraph_format.line_spacing=1.05
p0=DOC.add_paragraph();run=p0.add_run('VineWeave Canopy');run.bold=True;run.font.size=Pt(22);run.font.color.rgb=RGBColor.from_string('173C51');DOC.add_paragraph('green building | Group 12 | Integrated one-page design brief')
sections=[
('Purpose and application','Provide a local shaded pause point in an open campus, pocket park or comparable public pedestrian space. The objective is to reduce direct solar exposure while keeping circulation, root space and maintenance access clear; this is not a claim of street-wide air cooling.'),
('Geometry and parametric design','The tree-like, rounded shallow-arched support has 59 members, an overall extent of approximately 6.22 × 6.66 m and a maximum height of 3.48 m. The approximate maximum horizontal exterior span is 6.91 m. Ground-reaching support endpoint chords incline 45.9°–50.7° above horizontal; this is a defined geometric descriptor, not a local member angle.'),
('Materials, planting and construction','Twelve galvanised Q355B Ø80×4 mm support members are proposed with 47 solid Ø80 mm engineered-bamboo crown members. Product identity and properties require confirmation. The source planting proposal uses two Boston ivy plants and one Rangoon creeper per unit, with local suitability, attachment and winter behaviour to be checked. Prefabricated numbered members connect through proposed head nodes, sleeves and clamps. Grouting requires curing; complete dry assembly and lossless reversibility are not established.'),
('Environmental logic and analysis','The crown and foliage intercept direct sun; leaf placement and transmittance affect useful shade, while open sides retain potential airflow paths. Evapotranspiration depends on plant and water status. The 30°N example uses a 23.23 m² reference zone: ideal continuous-leaf solstice overlap is 47.2% / 98.7% / 47.2% at 09/12/15 solar time, with a 75.3% time mean. Seven parameter models and five spatial-cover states extend the analysis; these optical scenarios are neither plant-growth-duration predictions nor temperature reductions.'),
('Implementation and evidence boundary','The current unit budget is approximately CNY 51,676 and the carbon estimate approximately 2,820 kgCO₂e, both preliminary. Final product factors, joint/global behaviour, wind and foundations, quotations and the actual site require checking. Structural acceptance and vegetation shade readiness are separate milestones; no evidence-based maturity month is yet available. The next step is to validate occupied-hour shade and environmental response under site-specific conditions.')]
for title,body in sections:
 pp=DOC.add_paragraph();rr=pp.add_run(title);rr.bold=True;rr.font.color.rgb=RGBColor.from_string('346A55');DOC.add_paragraph(body)
DOC.add_paragraph('Basis: project model and take-off (P1), ideal shade analysis (P2), Cost-Analysis.xlsx (P3), design/build-up (P4), parameter and spatial-cover supplement (P5).')
DOC.save(O/'One-Page-Design-Brief.docx')
# Standalone editable selection/cost/logic workbook.
wb=Workbook();wb.remove(wb.active);wb.calculation=CalcProperties(fullCalcOnLoad=True)
def sheet(name,headers,rows,widths):
 s=wb.create_sheet(name);s.append(headers)
 for row in rows:s.append(row)
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.fill=PatternFill('solid',fgColor='346A55');c.font=Font(name='Calibri',bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True,vertical='center')
 for row in s.iter_rows(min_row=2):
  for c in row:c.font=Font(name='Calibri',size=11);c.alignment=Alignment(wrap_text=True,vertical='top')
 for k,w in widths.items():s.column_dimensions[k].width=w
 for i in range(2,s.max_row+1):s.row_dimensions[i].height=43
 return s
params=[['Overall height',3.48,'m','Rounded exterior dimension','Project model P1'],['Plan X / Y','6.22 / 6.66','m','Rounded exterior extents','Project model P1'],['Maximum horizontal span',6.91,'m','Max distance of projected exterior-mesh vertices','Supplemental model calculation'],['Support chord inclination','45.9–50.7','degrees','Above horizontal; endpoint chords, not local tangents','Supplemental model calculation'],['Member outside diameter',80,'mm','Geometric exterior diameter','Project model P1'],['Crown corner factor',.85,'—','Selected shape control','Project parameters P1'],['Upper expansion',1.08,'—','Selected shape control','Project parameters P1'],['Final member length',288.66,'m','Project take-off; 12 supports + 47 crown members','Project workbook P3']]
sheet('Final Parameters',['Parameter','Value','Unit','Definition / status','Source'],params,{'A':27,'B':22,'C':12,'D':64,'E':38})
source=load_workbook('/workspace/attachments/fc6700ae-cd4d-42be-9f81-e9aaa7e2f990/Cost-Analysis.xlsx',data_only=True);s=source['材料比选'];names=['Q355B solid Ø80','Q355B tube Ø80×3','Q355B tube Ø80×4','Q355B tube Ø80×5','Weathering steel Ø80×4','304 steel Ø80×3','Engineered bamboo solid Ø80','Treated timber solid Ø80','Raw bamboo Ø80'];rows=[]
for i,n in enumerate(names,5):rows.append([n,s.cell(i,8).value,s.cell(i,9).value,s.cell(i,10).value,s.cell(i,13).value,'Coating/exposure check'if i<=10 else'Treatment, moisture and replacement check','Reference workbook input; no verified product specification or quotation'])
sheet('Material Screen',['Material / section','Nominal bending kN·m','Conductivity W/(m·K)','Carbon factor kgCO2e/kg','Material cost CNY / 288.66m','Durability / exposure','Status'],rows,{'A':34,'B':24,'C':27,'D':29,'E':34,'F':49,'G':66})
cd=source['造价明细'];cost=[[cd.cell(i,3).value.split(' ')[-1]if False else n,cd.cell(i,4).value,'CNY','Seven-category estimate from P3; not a quotation']for i,n in zip(range(46,53),['Materials','Fabrication','Transport','Installation','Planting','Maintenance (3y)','Other'])]
cs=sheet('Cost Estimate',['Category','Amount','Unit','Basis / status'],cost,{'A':28,'B':22,'C':15,'D':75});cs.append(['TOTAL','=SUM(B2:B8)','CNY','Source total approximately 51,675.72; Excel recalculates on open.']);cs.cell(cs.max_row,2).number_format='#,##0.00'
for row in cs.iter_rows(min_row=2,max_col=2):row[1].number_format='#,##0.00'
logic=[['Broad crown + climbers','Intercept direct sun','Less exposure of people/ground','Geometric/optical calculation P2/P5; no measured cooling'],['Leaf cover and location','Transmission and overlap','Shade depends on amount and placement','Five cover states + distribution tests P5'],['Open sides / porosity','Potential paths and vegetation drag','Possible airflow trade-off','Qualitative design P4; no flow simulation'],['Healthy plants + water','Conditional evapotranspiration','Potential latent heat exchange','Literature mechanism; no project energy balance']]
sheet('Feature Cooling Logic',['Feature','Mechanism','Expected benefit','Evidence / condition'],logic,{'A':35,'B':42,'C':49,'D':76});wb.save(O/'Material-Parameter-and-Cost-Schedules.xlsx')
print('Prepared technical figures, two vector drawing sheets, group brief and editable schedules.')
